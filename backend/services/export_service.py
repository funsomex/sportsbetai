"""
Export service for generating Excel and PDF reports
"""
import io
from datetime import datetime, timezone
from typing import List, Dict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT


def generate_parlays_excel(parlays: List[Dict], stats: Dict) -> bytes:
    """Generate Excel file with parlays history and stats"""
    wb = Workbook()
    
    # Stats sheet
    ws_stats = wb.active
    ws_stats.title = "Estadísticas"
    
    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="CCFF00", end_color="CCFF00", fill_type="solid")
    header_font_dark = Font(bold=True, color="000000")
    
    # Stats section
    ws_stats["A1"] = "ESTADÍSTICAS DE APUESTAS"
    ws_stats["A1"].font = Font(bold=True, size=16)
    ws_stats.merge_cells("A1:D1")
    
    stats_data = [
        ["Métrica", "Valor"],
        ["Total Combinadas", stats.get("total", 0)],
        ["Ganadas", stats.get("won", 0)],
        ["Perdidas", stats.get("lost", 0)],
        ["Pendientes", stats.get("pending", 0)],
        ["Win Rate", f"{stats.get('win_rate', 0)}%"],
        ["ROI", f"{stats.get('roi', 0)}%"],
        ["Total Apostado", f"${stats.get('total_stake', 0)}"],
        ["Beneficio Total", f"${stats.get('total_profit', 0)}"],
        ["Mejor Ganancia", f"${stats.get('best_win', 0)}"],
        ["Peor Pérdida", f"${stats.get('worst_loss', 0)}"],
        ["Racha Actual", f"{stats.get('current_streak', 0)} {stats.get('streak_type', '')}"],
    ]
    
    for row_idx, row in enumerate(stats_data, start=3):
        for col_idx, value in enumerate(row, start=1):
            cell = ws_stats.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 3:  # Header row
                cell.font = header_font_dark
                cell.fill = header_fill
    
    # Adjust column widths
    ws_stats.column_dimensions["A"].width = 20
    ws_stats.column_dimensions["B"].width = 15
    
    # Parlays history sheet
    ws_parlays = wb.create_sheet("Historial Combinadas")
    
    headers = ["Fecha", "Nombre", "Selecciones", "Cuota Total", "Apuesta", "Ganancia Potencial", "Estado", "Beneficio"]
    for col, header in enumerate(headers, start=1):
        cell = ws_parlays.cell(row=1, column=col, value=header)
        cell.font = header_font_dark
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    for row_idx, parlay in enumerate(parlays, start=2):
        # Format selections
        selections_text = "\n".join([
            f"{s.get('match', 'N/A')}: {s.get('selection', 'N/A')}"
            for s in parlay.get("selections", [])
        ])
        
        ws_parlays.cell(row=row_idx, column=1, value=parlay.get("created_at", "")[:10])
        ws_parlays.cell(row=row_idx, column=2, value=parlay.get("name", ""))
        ws_parlays.cell(row=row_idx, column=3, value=selections_text)
        ws_parlays.cell(row=row_idx, column=4, value=parlay.get("total_odds", 0))
        ws_parlays.cell(row=row_idx, column=5, value=parlay.get("stake", 0) or 0)
        ws_parlays.cell(row=row_idx, column=6, value=parlay.get("potential_profit", 0) or 0)
        
        status = parlay.get("status", "pending")
        status_cell = ws_parlays.cell(row=row_idx, column=7, value=status.upper())
        if status == "won":
            status_cell.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
        elif status == "lost":
            status_cell.fill = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
        
        ws_parlays.cell(row=row_idx, column=8, value=parlay.get("actual_profit", 0) or 0)
    
    # Adjust column widths
    widths = [12, 20, 50, 12, 10, 15, 12, 12]
    for i, width in enumerate(widths, start=1):
        ws_parlays.column_dimensions[chr(64 + i)].width = width
    
    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generate_stats_pdf(stats: Dict, parlays: List[Dict]) -> bytes:
    """Generate PDF report with statistics"""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=50, bottomMargin=50)
    
    elements = []
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#CCFF00'),
        spaceAfter=30,
        alignment=TA_CENTER
    )
    
    subtitle_style = ParagraphStyle(
        'CustomSubtitle',
        parent=styles['Heading2'],
        fontSize=16,
        textColor=colors.HexColor('#333333'),
        spaceAfter=20,
        spaceBefore=20
    )
    
    # Title
    elements.append(Paragraph("SPORTSBETAI", title_style))
    elements.append(Paragraph("Reporte de Estadísticas", subtitle_style))
    elements.append(Spacer(1, 20))
    
    # Date
    date_style = ParagraphStyle('Date', parent=styles['Normal'], alignment=TA_CENTER)
    elements.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", date_style))
    elements.append(Spacer(1, 30))
    
    # Stats table
    stats_data = [
        ["MÉTRICA", "VALOR"],
        ["Total Combinadas", str(stats.get("total", 0))],
        ["Ganadas", str(stats.get("won", 0))],
        ["Perdidas", str(stats.get("lost", 0))],
        ["Pendientes", str(stats.get("pending", 0))],
        ["Win Rate", f"{stats.get('win_rate', 0)}%"],
        ["ROI", f"{stats.get('roi', 0)}%"],
        ["Total Apostado", f"${stats.get('total_stake', 0)}"],
        ["Beneficio Total", f"${stats.get('total_profit', 0)}"],
        ["Mejor Ganancia", f"${stats.get('best_win', 0)}"],
        ["Peor Pérdida", f"${stats.get('worst_loss', 0)}"],
    ]
    
    stats_table = Table(stats_data, colWidths=[200, 150])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#CCFF00')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F5F5F5')),
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#DDDDDD')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('TOPPADDING', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 8),
    ]))
    
    elements.append(stats_table)
    elements.append(Spacer(1, 40))
    
    # Recent parlays section
    elements.append(Paragraph("Últimas Combinadas", subtitle_style))
    
    recent_parlays = parlays[:10]  # Last 10
    
    if recent_parlays:
        parlay_data = [["Fecha", "Cuota", "Apuesta", "Estado", "Beneficio"]]
        
        for p in recent_parlays:
            status = p.get("status", "pending").upper()
            profit = p.get("actual_profit", 0) or 0
            parlay_data.append([
                p.get("created_at", "")[:10],
                f"{p.get('total_odds', 0):.2f}",
                f"${p.get('stake', 0) or 0}",
                status,
                f"${profit:.2f}"
            ])
        
        parlay_table = Table(parlay_data, colWidths=[80, 60, 80, 80, 80])
        parlay_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#333333')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9F9F9')]),
        ]))
        
        elements.append(parlay_table)
    else:
        elements.append(Paragraph("No hay combinadas registradas.", styles['Normal']))
    
    # Build PDF
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
